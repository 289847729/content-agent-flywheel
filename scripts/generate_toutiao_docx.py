import os
import docx
from docx import Document
from docx.shared import Pt, Inches, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml import OxmlElement, parse_xml
from docx.oxml.ns import qn, nsdecls

def create_toutiao_docx():
    doc = Document()

    # 设置页边距
    sections = doc.sections
    for section in sections:
        section.top_margin = Inches(0.8)
        section.bottom_margin = Inches(0.8)
        section.left_margin = Inches(0.85)
        section.right_margin = Inches(0.85)

    # 颜色定义
    COLOR_PRIMARY = RGBColor(226, 44, 56)     # 头条红
    COLOR_TITLE = RGBColor(34, 34, 34)        # 深黑标题
    COLOR_TEXT = RGBColor(51, 51, 51)         # 正文灰黑
    COLOR_MUTED = RGBColor(120, 120, 120)     # 辅助灰
    COLOR_HEADER_BG = "F4F5F7"

    # 设置默认正文字体
    normal_style = doc.styles['Normal']
    normal_style.font.name = 'PingFang SC'
    normal_style.font.size = Pt(11)
    normal_style.font.color.rgb = COLOR_TEXT
    normal_style._element.rPr.rFonts.set(qn('w:eastAsia'), 'PingFang SC')

    def set_run_font(run, font_name='PingFang SC', size_pt=11, bold=False, color_rgb=COLOR_TEXT, italic=False):
        run.font.name = font_name
        run.font.size = Pt(size_pt)
        run.bold = bold
        run.italic = italic
        run.font.color.rgb = color_rgb
        run._element.rPr.rFonts.set(qn('w:eastAsia'), font_name)

    def add_p(text="", align=WD_ALIGN_PARAGRAPH.LEFT, space_before=0, space_after=6, line_spacing=1.35):
        p = doc.add_paragraph()
        p.alignment = align
        p.paragraph_format.space_before = Pt(space_before)
        p.paragraph_format.space_after = Pt(space_after)
        p.paragraph_format.line_spacing = line_spacing
        return p

    def add_h2(text):
        p = add_p(space_before=16, space_after=8)
        # 装饰红标
        run_tag = p.add_run("▌ ")
        set_run_font(run_tag, size_pt=14, bold=True, color_rgb=COLOR_PRIMARY)
        run_text = p.add_run(text)
        set_run_font(run_text, size_pt=14, bold=True, color_rgb=COLOR_TITLE)
        return p

    # --- 1. 文章主标题 ---
    selected_title = "全球最强开源大模型易主！雷军悄悄干到世界第一，连训练过程都全网直播了？"
    p_title = add_p(space_before=4, space_after=12, line_spacing=1.2)
    run_t = p_title.add_run(selected_title)
    set_run_font(run_t, size_pt=18, bold=True, color_rgb=COLOR_TITLE)

    # --- 2. 插入高清封面图片 ---
    cover_image_path = "/Users/mrfan/dev/ai-dev/wechat-content-flywheel/outputs/2026-09-23-小米-MiMo-V2.6-今日头条封面.jpg"
    if os.path.exists(cover_image_path):
        p_img = add_p(align=WD_ALIGN_PARAGRAPH.CENTER, space_before=4, space_after=14)
        run_img = p_img.add_run()
        # 宽度适合页面宽度 (可用宽度约 8.5 - 2*0.85 = 6.8 Inches)
        run_img.add_picture(cover_image_path, width=Inches(6.6))

    # 分割线
    p_hr = add_p(space_after=12)
    p_hr_border = parse_xml(r'<w:pBdr xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main">'
                            r'<w:bottom w:val="single" w:sz="6" w:space="1" w:color="E5E5E5"/>'
                            r'</w:pBdr>')
    p_hr._p.get_or_add_pPr().append(p_hr_border)

    # --- 3. 导语 ---
    p = add_p(space_after=8)
    r1 = p.add_run("这两天科技圈彻底炸锅了，所有人的目光都聚焦在同一件事上：")
    set_run_font(r1)
    r2 = p.add_run("小米突然发布了自研大模型 MiMo-V2.6，并在国际开源权威榜单上一举登顶！")
    set_run_font(r2, bold=True, color_rgb=COLOR_PRIMARY)

    p = add_p(space_after=8)
    r1 = p.add_run("很多人看到第一反应可能跟我一样：")
    set_run_font(r1)
    r2 = p.add_run("小米不是做手机、做家电、造汽车的吗？怎么不声不响，在大模型跑分上把一众 AI 专业大厂给反超了？")
    set_run_font(r2, bold=True)

    p = add_p(space_after=8)
    r1 = p.add_run("更夸张的是，这次小米不仅放出了万亿参数的大模型，还把训练过程的“底裤”彻底掀开全网直播。")
    set_run_font(r1)

    p = add_p(space_after=12)
    r1 = p.add_run("今天咱们不讲虚头巴脑的套话，用人话拆一拆：")
    set_run_font(r1)
    r2 = p.add_run("小米这次到底有多猛？和闭源顶尖模型相比差距在哪？以及我们普通人到底能用上什么？")
    set_run_font(r2, bold=True)

    # --- 4. 一、开源榜单突然易主 ---
    add_h2("一、开源榜单突然易主：小米凭什么拿下全球第一？")

    p = add_p(space_after=8)
    r1 = p.add_run("这次引起巨大轰动的数据，来自国际公认的 ")
    set_run_font(r1)
    r2 = p.add_run("Artificial Analysis 综合智能指数")
    set_run_font(r2, bold=True)
    r3 = p.add_run("。根据最新榜单显示，小米 ")
    set_run_font(r3)
    r4 = p.add_run("MiMo-V2.6-Pro")
    set_run_font(r4, bold=True)
    r5 = p.add_run(" 拿下了 ")
    set_run_font(r5)
    r6 = p.add_run("46 分")
    set_run_font(r6, bold=True, color_rgb=COLOR_PRIMARY)
    r7 = p.add_run(" 的高分，正式超越了此前长期霸榜的熟面孔，拿下全球开源大模型第一位。")
    set_run_font(r7)

    p = add_p(space_after=8)
    r1 = p.add_run("大家可以直观感受一下这份最新的开源头部排名：")
    set_run_font(r1)

    # 表格 1
    table1 = doc.add_table(rows=5, cols=3)
    table1.alignment = WD_TABLE_ALIGNMENT.CENTER
    headers = ["排名与模型", "智能指数得分", "状态"]
    data1 = [
        ["MiMo-V2.6-Pro（小米）", "46 分（全球第一）", "开源"],
        ["GLM-5.3", "45 分", "开源"],
        ["Kimi K3", "44 分", "开源"],
        ["MiMo-V2.5-Pro（上一代）", "26 分", "开源"]
    ]

    for col_idx, h in enumerate(headers):
        cell = table1.cell(0, col_idx)
        cell.text = h
        shading = parse_xml(r'<w:shd xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main" w:fill="F4F5F7"/>')
        cell._tc.get_or_add_tcPr().append(shading)
        p_cell = cell.paragraphs[0]
        p_cell.alignment = WD_ALIGN_PARAGRAPH.CENTER
        for r in p_cell.runs:
            set_run_font(r, size_pt=10, bold=True, color_rgb=COLOR_TITLE)

    for row_idx, row_data in enumerate(data1):
        for col_idx, val in enumerate(row_data):
            cell = table1.cell(row_idx + 1, col_idx)
            cell.text = val
            p_cell = cell.paragraphs[0]
            if col_idx == 0:
                p_cell.alignment = WD_ALIGN_PARAGRAPH.LEFT
            else:
                p_cell.alignment = WD_ALIGN_PARAGRAPH.CENTER
            for r in p_cell.runs:
                is_mimo = (row_idx == 0)
                set_run_font(r, size_pt=10, bold=is_mimo, color_rgb=COLOR_PRIMARY if (is_mimo and col_idx==1) else COLOR_TEXT)

    # 设置表格边框与内边距
    tblBorders = parse_xml(
        r'<w:tblBorders xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main">'
        r'<w:top w:val="single" w:sz="4" w:space="0" w:color="E0E0E0"/>'
        r'<w:bottom w:val="single" w:sz="6" w:space="0" w:color="B0B0B0"/>'
        r'<w:left w:val="none"/>'
        r'<w:right w:val="none"/>'
        r'<w:insideH w:val="single" w:sz="4" w:space="0" w:color="EEEEEE"/>'
        r'<w:insideV w:val="none"/>'
        r'</w:tblBorders>'
    )
    table1._tbl.tblPr.append(tblBorders)

    # 表格后空行
    add_p(space_after=6)

    p = add_p(space_after=8)
    r1 = p.add_run("最让人震惊的是前后代跨度对比：就在不到一年前，小米的 ")
    set_run_font(r1)
    r2 = p.add_run("上一代 V2.5-Pro 还只有 26 分")
    set_run_font(r2, bold=True)
    r3 = p.add_run("。短短几个月内，直接")
    set_run_font(r3)
    r4 = p.add_run("暴涨了整整 20 分")
    set_run_font(r4, bold=True, color_rgb=COLOR_PRIMARY)
    r5 = p.add_run("！在大模型卷到今天这个瓶颈期，每一分的提升都极其艰难，这种跨越式提升确实打破了很多人的常识。")
    set_run_font(r5)

    p = add_p(space_after=12)
    r1 = p.add_run("不过咱们客观来看，这次登顶有个核心技术细节：")
    set_run_font(r1)
    r2 = p.add_run("MiMo-V2.6-Pro 走的是 MoE（混合专家）架构，总参数高达 1.02 万亿，但实际激活参数只有 420 亿")
    set_run_font(r2, bold=True)
    r3 = p.add_run("。也就是说，平时虽然体量庞大，但干活时只调用对应领域的“专家模块”，既保证了脑容量，又极大控制了运行消耗。")
    set_run_font(r3)

    # --- 5. 二、罕见“掀开底裤” ---
    add_h2("二、罕见“掀开底裤”：训练全程直播，菜谱全都送给你")

    p = add_p(space_after=8)
    r1 = p.add_run("如果说拿第一只是成绩，那这次更让整个技术圈震动的，是小米的")
    set_run_font(r1)
    r2 = p.add_run("打法")
    set_run_font(r2, bold=True)
    r3 = p.add_run("。")
    set_run_font(r3)

    p = add_p(space_after=8)
    r1 = p.add_run("过去国内外大厂做开源，基本上是自己在机房里悄悄跑好几个月，跑出结果后把最终模型权重打包丢到网上，整个训练过程完全是个“黑盒子”。而这次小米做了件行业里几乎没人敢做的事：")
    set_run_font(r1)
    r2 = p.add_run("直接把训练全程搬到网上直播！")
    set_run_font(r2, bold=True, color_rgb=COLOR_PRIMARY)

    p = add_p(space_after=6)
    r1 = p.add_run("这几个硬核数据公布出来，直接让不少开发者看愣了：")
    set_run_font(r1)

    bullets = [
        ("30 个强化学习（RL）step", "：从头到尾全流程透明展示；"),
        ("不到 6 天时间", "：从正式开训到成型发布，仅用不到一周；"),
        ("约 350 万美元", "：消耗总训练成本，性价比极高；"),
        ("整整 75 万条", "：同步公开全部高质量推理训练轨迹；"),
        ("不到 1 块钱", "：单个复杂任务的折算运行成本，极度亲民。")
    ]
    for b_title, b_desc in bullets:
        pb = add_p(space_after=4)
        pb.paragraph_format.left_indent = Inches(0.2)
        rb_dot = pb.add_run("●  ")
        set_run_font(rb_dot, size_pt=9, bold=True, color_rgb=COLOR_PRIMARY)
        rb_title = pb.add_run(b_title)
        set_run_font(rb_title, bold=True)
        rb_desc = pb.add_run(b_desc)
        set_run_font(rb_desc)

    p = add_p(space_before=6, space_after=8)
    r1 = p.add_run("这相当于什么？相当于一家顶级名厨不仅免费请你吃饭，还把祖传菜谱、火候时间、甚至每道食材的进货价和供货商全部贴在大门口。")
    set_run_font(r1)

    p = add_p(space_after=12)
    r1 = p.add_run("这事对国内开源生态的影响，可能远超跑分本身。中小团队和独立开发者再想复现或者微调，不用再两眼一抹黑花天价去试错，直接照着小米公开的轨迹走就行，")
    set_run_font(r1)
    r2 = p.add_run("整个行业的试错成本被狠狠打了下来")
    set_run_font(r2, bold=True)
    r3 = p.add_run("。")
    set_run_font(r3)

    # --- 6. 三、普通人别看热闹 ---
    add_h2("三、普通人别看热闹：这玩意儿装在哪能用？")

    p = add_p(space_after=8)
    r1 = p.add_run("聊完技术，大家最关心的肯定还是：")
    set_run_font(r1)
    r2 = p.add_run("这东西再厉害，跟我一个普通打工人、日常手机用户有什么关系？")
    set_run_font(r2, bold=True)

    p = add_p(space_after=8)
    r1 = p.add_run("这次小米并没有只给程序员画饼，而是同步推出了几款普通人能“零门槛直接用”的接地气工具：")
    set_run_font(r1)

    features = [
        ("MiMo Desktop 电脑客户端", "不用去网页端注册排队，直接下载客户端到本地电脑就能跑。最大的好处是本地离线处理，企业和个人的私密资料完全不用担心上传泄露。"),
        ("Flash 轻量版同步开源", "如果你的电脑显卡不是顶配，官方提供了 3090 亿参数的轻量版，大幅降低硬件门槛。"),
        ("100 万 Token 超长上下文", "相当于一整部长篇小说、一本厚厚的行业年报，你可以一次性直接扔进去，让它几秒钟出摘要、找漏洞。"),
        ("UltraSpeed 20 倍速模式", "大幅解决了以往万亿大模型打字像“挤牙膏”一样的痛点，文字生成飞快。"),
        ("真正的原生全模态", "不仅能打字聊天，还能直接看懂复杂工程图纸、听懂长音频会议录音、甚至理解多段视频。")
    ]
    for idx, (f_title, f_desc) in enumerate(features, 1):
        pf = add_p(space_after=5)
        pf.paragraph_format.left_indent = Inches(0.2)
        rf_num = pf.add_run(f"{idx}.  ")
        set_run_font(rf_num, bold=True, color_rgb=COLOR_PRIMARY)
        rf_title = pf.add_run(f"【{f_title}】 ")
        set_run_font(rf_title, bold=True)
        rf_desc = pf.add_run(f_desc)
        set_run_font(rf_desc)

    p = add_p(space_before=6, space_after=12)
    r1 = p.add_run("日常场景里，比如甩给它一份 50 页的竞品 PDF 问核心结论、丢一段 2 小时的杂乱会议录音出纪要，都能直接派上用场。")
    set_run_font(r1)

    # --- 7. 四、理性看待 ---
    add_h2("四、理性看待：先别急着吹“世界无敌”")

    p = add_p(space_after=8)
    r1 = p.add_run("最后，咱们必须泼一点理性的冷水，不搞盲目自嗨。虽然小米在")
    set_run_font(r1)
    r2 = p.add_run("开源领域")
    set_run_font(r2, bold=True)
    r3 = p.add_run("拿下了全球第一，但如果放眼全球整个 AI 领域，")
    set_run_font(r3)
    r4 = p.add_run("开源与闭源头部旗舰依然存在肉眼可见的差距")
    set_run_font(r4, bold=True, color_rgb=COLOR_PRIMARY)
    r5 = p.add_run("：")
    set_run_font(r5)

    # 表格 2
    table2 = doc.add_table(rows=4, cols=3)
    table2.alignment = WD_TABLE_ALIGNMENT.CENTER
    headers2 = ["模型", "综合智能分数", "阵营类别"]
    data2 = [
        ["GPT-6 Astra", "53 分", "闭源（领先）"],
        ["Claude Fable 5.1", "53 分", "闭源（领先）"],
        ["MiMo-V2.6-Pro（小米）", "46 分", "开源（全球第一）"]
    ]

    for col_idx, h in enumerate(headers2):
        cell = table2.cell(0, col_idx)
        cell.text = h
        shading = parse_xml(r'<w:shd xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main" w:fill="F4F5F7"/>')
        cell._tc.get_or_add_tcPr().append(shading)
        p_cell = cell.paragraphs[0]
        p_cell.alignment = WD_ALIGN_PARAGRAPH.CENTER
        for r in p_cell.runs:
            set_run_font(r, size_pt=10, bold=True, color_rgb=COLOR_TITLE)

    for row_idx, row_data in enumerate(data2):
        for col_idx, val in enumerate(row_data):
            cell = table2.cell(row_idx + 1, col_idx)
            cell.text = val
            p_cell = cell.paragraphs[0]
            if col_idx == 0:
                p_cell.alignment = WD_ALIGN_PARAGRAPH.LEFT
            else:
                p_cell.alignment = WD_ALIGN_PARAGRAPH.CENTER
            for r in p_cell.runs:
                is_mimo = (row_idx == 2)
                set_run_font(r, size_pt=10, bold=is_mimo, color_rgb=COLOR_PRIMARY if (is_mimo and col_idx==1) else COLOR_TEXT)

    table2._tbl.tblPr.append(tblBorders)
    add_p(space_after=6)

    p = add_p(space_after=8)
    r1 = p.add_run("在 50 分以上的神仙打架区间，")
    set_run_font(r1)
    r2 = p.add_run("7 分的差距绝对不容忽视")
    set_run_font(r2, bold=True)
    r3 = p.add_run("。在最顶级的复杂逻辑推理和极限创作上，海外顶尖闭源模型依然处于压制地位。")
    set_run_font(r3)

    p = add_p(space_after=8)
    r1 = p.add_run("所以，客观准确的结论应该是：")
    set_run_font(r1)
    r2 = p.add_run("在开源这条赛道上，小米暂时跑到了全球最前列，但离世界真正的终极旗舰，还有追赶空间。")
    set_run_font(r2, bold=True)

    p = add_p(space_after=14)
    r1 = p.add_run("不过对于我们普通人来说，这绝对是天大的利好。大厂开源卷得越狠、技术越透明，最终好用的 AI 工具就会越来越便宜，甚至彻底免费。")
    set_run_font(r1)

    # --- 8. 互动探讨（头条评论引导框） ---
    callout_p = add_p(space_before=10, space_after=10)
    callout_p.paragraph_format.left_indent = Inches(0.2)
    callout_p.paragraph_format.right_indent = Inches(0.2)

    callout_border = parse_xml(r'<w:pBdr xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main">'
                               r'<w:left w:val="single" w:sz="24" w:space="8" w:color="E22C38"/>'
                               r'</w:pBdr>')
    callout_shd = parse_xml(r'<w:shd xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main" w:fill="FAF0F0"/>')
    callout_p._p.get_or_add_pPr().append(callout_border)
    callout_p._p.get_or_add_pPr().append(callout_shd)

    r_talk_title = callout_p.add_run("【互动探讨 · 头条热议】\n")
    set_run_font(r_talk_title, size_pt=10.5, bold=True, color_rgb=COLOR_PRIMARY)
    r_talk_content = callout_p.add_run("你觉得小米这次大模型登顶是厚积薄发还是营销造势？如果电脑上能免费用，你最希望它帮你干什么活？欢迎在评论区留下你的看法，一起聊聊！")
    set_run_font(r_talk_content, size_pt=10.5, bold=True, color_rgb=COLOR_TITLE)

    # 导出保存
    output_path = "/Users/mrfan/dev/ai-dev/wechat-content-flywheel/outputs/2026-09-23-小米-MiMo-V2.6-今日头条文案.docx"
    doc.save(output_path)
    try:
        print(f"Successfully generated docx with cover image at: {output_path}")
    except OSError:
        pass

if __name__ == "__main__":
    create_toutiao_docx()
